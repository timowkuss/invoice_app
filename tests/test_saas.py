import io
import os
from decimal import Decimal
import pytest
from PIL import Image
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from core.matching.service import fuzzy_score, extract_numbers, MatchingService
from core.auth import create_access_token, decode_access_token
from models.product import Product as OcrProduct

@pytest.fixture
def client(tmp_path,monkeypatch):
    import api.main as main
    import api.routers.documents as docs
    import core.uploads as uploads
    monkeypatch.setattr(main,'DATABASE_URL',f'sqlite:///{tmp_path / "test.db"}')
    monkeypatch.setattr(docs,'UPLOAD_DIR',tmp_path/'uploads')
    monkeypatch.setattr(uploads,'UPLOAD_DIR',tmp_path/'uploads')
    monkeypatch.setenv('SECRET_KEY','test-secret-'*5)
    monkeypatch.setenv('OCR_WORKER','false')
    monkeypatch.setenv('MISTRAL_API_KEY','test-only')
    monkeypatch.setenv('ADMIN_EMAIL','owner@example.com')
    monkeypatch.setenv('ADMIN_PASSWORD','Owner-password-1234')
    with TestClient(main.app) as c:
        yield c

def signup(c,name='shop_a'):
    owner=c.post('/api/auth/login',json={'email':'owner@example.com','password':'Owner-password-1234'})
    assert owner.status_code==200,owner.text
    admin=owner.json()['user']
    ah={'Authorization':'Bearer '+create_access_token(admin['id'],admin['role'],admin['store_id'])}
    store=c.post('/api/stores/',headers=ah,json={'name':name}).json()
    email=f'{name}@example.com'
    response=c.post('/api/auth/register',headers=ah,json={'email':email,'password':'Test-password-1234','full_name':name,'role':'store_admin','store_id':store['id']})
    assert response.status_code==201,response.text
    u=response.json()
    login=c.post('/api/auth/login',json={'email':email,'password':'Test-password-1234'})
    assert login.status_code==200,login.text
    return u,{'Authorization':'Bearer '+create_access_token(u['id'],u['role'],u['store_id'])}

def catalog(c,headers,rows=None):
    text='Код 1С;Название;Штрихкод;Единица;Цена\n'+(rows or '000001;Шадринское Молоко 1л 5%;0012345678901;шт;100\n000002;Шадринское Молоко 1л 3.2%;0012345678902;шт;95\n')
    result=c.post('/api/products/import',headers=headers,files={'file':('catalog.csv',text.encode('utf-8'),'text/csv')})
    assert result.status_code==200,result.text
    return c.get('/api/products/',headers=headers).json()['items']

def upload(c,headers,color='white'):
    file=io.BytesIO();Image.new('RGB',(100,100),color).save(file,format='PNG')
    result=c.post('/api/documents/upload',headers=headers,files={'file':('invoice.png',file.getvalue(),'image/png')})
    assert result.status_code==201,result.text
    return result.json()

def process(c,headers,doc,monkeypatch,products=None):
    import core.worker as worker
    from core.repositories import DocumentRepo
    async def fake(path):
        return (products or [OcrProduct(name='Молоко Шадринское 5% 1л',quantity='2',price='100,00',total='200,00',unit='шт')]),1
    monkeypatch.setattr(worker,'recognize',fake)
    response=c.post(f'/api/documents/{doc["id"]}/process',headers=headers)
    assert response.status_code==202,response.text
    async def run():
        await DocumentRepo.update_status(doc['id'],'processing')
        await worker.process_claimed(await DocumentRepo.get_by_id(doc['id']))
    c.portal.call(run)
    result=c.get(f'/api/documents/{doc["id"]}',headers=headers).json()
    assert result['document']['status']=='needs_review',result
    return result

def test_matching_order_and_variants():
    assert fuzzy_score('Молоко Шадринское 5% 1л','Шадринское Молоко 1л 5%')==100
    assert fuzzy_score('Молоко Шадринское 5% 1л','Шадринское Молоко 1л 3.2%')==0
    assert fuzzy_score('Молоко Шадринское 5% 1л','Шадринское Молоко 0.5л 5%')==0
    assert fuzzy_score('Шадринское Молоко','Шадринское Молоко 1л 5%')<92
    assert extract_numbers('500мл 250г 3,2%')=={'volume':.5,'weight':.25,'fat_content':3.2}

def test_complete_flow_and_excel(client,monkeypatch):
    _,h=signup(client);catalog(client,h);doc=upload(client,h)
    result=process(client,h,doc,monkeypatch)
    item=result['items'][0]
    assert item['product_name']=='Шадринское Молоко 1л 5%'
    assert item['match_status']=='matched'
    assert item['total']==200
    url=f'/api/documents/{doc["id"]}'
    assert client.get(url+'/export',headers=h).status_code==409
    assert client.post(url+'/confirm',headers=h).status_code==200
    response=client.get(url+'/export',headers=h)
    assert response.status_code==200,response.text
    book=load_workbook(io.BytesIO(response.content));sheet=book.active
    assert sheet['A2'].value=='000001'
    assert sheet['B2'].value=='Шадринское Молоко 1л 5%'
    assert sheet['D2'].value=='0012345678901'
    assert sheet['E2'].value==2 and sheet['G2'].value==200
    assert sheet['D2'].data_type=='s'
    assert client.get('/api/stats/dashboard',headers=h).json()['confirmed']==1
    usage=client.get('/api/stats/ai-usage',headers=h).json()
    assert usage['stats']['total_requests']==1 and usage['stats']['total_pages']==1
    assert usage['recent'][0]['pages']==1

def test_super_admin_ai_usage_across_stores(client,monkeypatch):
    monkeypatch.setenv('MISTRAL_COST_PER_1000_PAGES_USD','2.5')
    _,ha=signup(client);process(client,ha,upload(client,ha),monkeypatch)
    _,hb=signup(client,'shop_b');process(client,hb,upload(client,hb,'red'),monkeypatch)
    assert client.get('/api/stats/admin-ai-usage',headers=ha).status_code==403
    owner=client.post('/api/auth/login',json={'email':'owner@example.com','password':'Owner-password-1234'}).json()['user']
    headers={'Authorization':'Bearer '+create_access_token(owner['id'],owner['role'],owner['store_id'])}
    response=client.get('/api/stats/admin-ai-usage',headers=headers)
    assert response.status_code==200,response.text
    data=response.json()
    assert data['summary']['total_requests']==2
    assert data['summary']['successful_requests']==2
    assert data['summary']['failed_requests']==0
    assert data['summary']['total_pages']==2
    assert data['summary']['total_cost']==pytest.approx(0.005)
    assert {row['store_name'] for row in data['stores']}=={'shop_a','shop_b'}
    assert len(data['recent'])==2

def test_tenant_isolation_and_role_escalation(client,monkeypatch):
    a,ha=signup(client);pa=catalog(client,ha);doc=upload(client,ha)
    result=process(client,ha,doc,monkeypatch)
    b,hb=signup(client,'shop_b');pb=catalog(client,hb)
    url=f'/api/documents/{doc["id"]}'
    for path in (url,url+'/image',url+'/audit',url+'/export',f'/api/products/{pa[0]["id"]}',f'/api/products/{pa[0]["id"]}/aliases'):
        assert client.get(path,headers=hb).status_code==404,path
    assert client.post(url+'/process',headers=hb).status_code==404
    assert client.post(url+'/confirm',headers=hb).status_code==404
    assert client.get('/api/documents/',headers=hb).json()['total']==0
    bad=client.put(url+f'/items/{result["items"][0]["id"]}',headers=ha,json={'product_id':pb[0]['id'],'create_alias':True})
    assert bad.status_code==404
    for payload in ({'role':'super_admin'},{'store_id':b['store_id']}):
        assert client.put(f'/api/users/{a["id"]}',headers=ha,json=payload).status_code==403
    assert client.post('/api/auth/register',headers=ha,json={'email':'hacker@example.com','password':'Test-password-1234','role':'store_admin','store_id':a['store_id']}).status_code==403
    assert client.post(f'/api/users/{b["id"]}/bind-telegram?chat_id=123',headers=ha).status_code==403
    assert 'test-only' not in client.get('/api/settings/',headers=ha).text
    assert client.put('/api/settings/',headers=ha,json={'key':'mistral_api_key','value':'x'}).status_code==405

def test_review_alias_and_edit_invalidates_confirmation(client,monkeypatch):
    _,h=signup(client);products=catalog(client,h);doc=upload(client,h)
    result=process(client,h,doc,monkeypatch,[OcrProduct(name='Молоко Шадр.',quantity='1',price='100',total='100',unit='шт')])
    url=f'/api/documents/{doc["id"]}';item=result['items'][0]
    assert client.post(url+'/confirm',headers=h).status_code==409
    response=client.put(url+f'/items/{item["id"]}',headers=h,json={'product_id':products[0]['id'],'quantity':'3','price':'0','create_alias':True})
    assert response.status_code==200,response.text
    assert response.json()['total']==0
    assert client.post(url+'/confirm',headers=h).status_code==200
    assert client.get(url+'/export',headers=h).status_code==200
    assert client.put(url+f'/items/{item["id"]}',headers=h,json={'quantity':'4'}).status_code==200
    assert client.get(url+'/export',headers=h).status_code==409
    assert len(client.get('/api/matching/aliases',headers=h).json())==1
    assert client.put(url+f'/items/{item["id"]}',headers=h,json={'quantity':'-1'}).status_code==422
    assert client.put(url+f'/items/{item["id"]}',headers=h,json={'price':'NaN'}).status_code==422

def test_upload_validation_duplicates_and_atomic_catalog(client):
    _,h=signup(client);upload(client,h)
    data=io.BytesIO();Image.new('RGB',(100,100),'white').save(data,format='PNG')
    assert client.post('/api/documents/upload',headers=h,files={'file':('invoice.png',data.getvalue())}).status_code==409
    for name,content,status in [('bad.exe',b'x',415),('bad.pdf',b'bad',422),('bad.jpg',data.getvalue(),422)]:
        assert client.post('/api/documents/upload',headers=h,files={'file':(name,content)}).status_code==status
    response=client.post('/api/products/import',headers=h,files={'file':('bad.csv','Код 1С;Название\n1;Товар\n1;Другой'.encode())})
    assert response.status_code==422
    assert client.get('/api/products/',headers=h).json()['total']==0

def test_authentication_and_csrf(client):
    u,h=signup(client)
    cookie=client.cookies.get('invoice_session')
    assert decode_access_token(cookie)['sub']==str(u['id'])
    assert client.post('/api/auth/logout',headers={'Origin':'https://evil.example'}).status_code==403
    assert client.get('/api/documents/',headers={'Authorization':'Bearer invalid'}).status_code==401
    assert client.post('/api/auth/logout').status_code==200
    assert client.get('/api/auth/me').status_code==401
    response=client.post('/api/auth/login',json={'email':'shop_a@example.com','password':'Test-password-1234'})
    assert response.status_code==200
    assert 'HttpOnly' in response.headers['set-cookie'] and 'SameSite=strict' in response.headers['set-cookie']
    assert 'token' not in response.json()
    assert client.post('/api/auth/login',json={'email':'shop_a@example.com','password':'wrong'}).status_code==401
    assert client.post('/api/auth/signup',json={'email':'new@example.com','password':'Test-password-1234'}).status_code==404

def test_ambiguous_match_does_not_autoselect(client):
    u,h=signup(client)
    catalog(client,h,'1;Шадринское Молоко 1л 5%;111;шт;100\n2;Шадринское Молоко 1л 5%;222;шт;100\n')
    result=client.portal.call(MatchingService().match_product,u['store_id'],'Молоко Шадринское 5% 1л')
    assert result.confidence<92
    assert len(result.alternatives)==2

def test_operator_cannot_import_or_create_accounts(client):
    u,h=signup(client)
    owner=client.post('/api/auth/login',json={'email':'owner@example.com','password':'Owner-password-1234'}).json()['user']
    ah={'Authorization':'Bearer '+create_access_token(owner['id'],owner['role'],u['store_id'])}
    response=client.post('/api/auth/register',headers=ah,json={'email':'operator@example.com','password':'Test-password-1234','role':'operator','store_id':u['store_id']})
    assert response.status_code==201,response.text
    op=response.json();oh={'Authorization':'Bearer '+create_access_token(op['id'],op['role'],op['store_id'])}
    assert client.post('/api/products/import',headers=oh,files={'file':('a.csv',b'x')}).status_code==403
    payload={'email':'other@example.com','password':'Test-password-1234','role':'operator','store_id':u['store_id']}
    assert client.post('/api/auth/register',headers=h,json=payload).status_code==403
    assert client.post('/api/auth/register',headers=oh,json=payload).status_code==403

def test_password_reset_by_email(client,monkeypatch):
    signup(client)
    sent={}
    monkeypatch.setattr('core.mailer.is_configured',lambda:True)
    def capture(email,url):
        sent.update(email=email,url=url)
    monkeypatch.setattr('core.mailer.send_password_reset_email',capture)
    response=client.post('/api/auth/forgot-password',json={'email':'shop_a@example.com'})
    assert response.status_code==200,response.text
    assert sent['email']=='shop_a@example.com'
    token=sent['url'].split('reset_token=',1)[1]
    response=client.post('/api/auth/reset-password',json={'token':token,'password':'New-password-1234'})
    assert response.status_code==200,response.text
    assert client.post('/api/auth/login',json={'email':'shop_a@example.com','password':'New-password-1234'}).status_code==200
    assert client.post('/api/auth/reset-password',json={'token':token,'password':'Another-password-1234'}).status_code==400

def test_ocr_failure_sanitized_and_retry_idempotent(client,monkeypatch):
    import core.worker as worker
    from core.repositories import DocumentRepo
    _,h=signup(client);doc=upload(client,h);url=f'/api/documents/{doc["id"]}'
    assert client.post(url+'/process',headers=h).status_code==202
    assert client.post(url+'/process',headers=h).status_code==202
    assert len([x for x in client.get(url+'/audit',headers=h).json() if x['action']=='queue_ocr'])==1
    async def bad(path):raise RuntimeError('SECRET MUST NOT LEAK')
    monkeypatch.setattr(worker,'recognize',bad)
    async def run():await worker.process_claimed(await DocumentRepo.get_by_id(doc['id']))
    client.portal.call(run)
    result=client.get(url,headers=h)
    assert result.json()['document']['status']=='error'
    assert 'SECRET MUST NOT LEAK' not in result.text

def test_add_delete_items_and_export(client,monkeypatch):
    _,h=signup(client);products=catalog(client,h);doc=upload(client,h)
    result=process(client,h,doc,monkeypatch)
    url=f'/api/documents/{doc["id"]}'
    assert client.post(url+'/confirm',headers=h).status_code==200
    created=client.post(url+'/items',headers=h,json={'product_id':products[1]['id'],'quantity':'1.125','price':'2.12'})
    assert created.status_code==201,created.text
    item=created.json()
    assert item['row_number']==2 and item['total']==2.39 and item['match_status']=='manual'
    assert item['ocr_text']=='' and item['product_name']==products[1]['name']
    detail=client.get(url,headers=h).json()['document']
    assert detail['status']=='needs_review' and detail['confirmed_at'] is None
    assert detail['total_items']==2 and detail['matched_items']==2
    assert client.get(url+'/export',headers=h).status_code==409
    assert client.post(url+'/confirm',headers=h).status_code==200
    old_id=result['items'][0]['id']
    assert client.delete(url+f'/items/{old_id}',headers=h).status_code==200
    detail=client.get(url,headers=h).json()
    assert detail['items'][0]['id']==item['id'] and detail['items'][0]['row_number']==1
    assert detail['document']['total_items']==1 and detail['document']['confirmed_at'] is None
    assert client.get(url+'/export',headers=h).status_code==409
    assert client.post(url+'/confirm',headers=h).status_code==200
    book=load_workbook(io.BytesIO(client.get(url+'/export',headers=h).content))
    assert book.active.max_row==2 and book.active['B2'].value==products[1]['name']
    assert book.active['G2'].value==2.39
    assert client.delete(url+f'/items/{item["id"]}',headers=h).status_code==200
    detail=client.get(url,headers=h).json()
    assert detail['items']==[] and detail['document']['total_items']==0
    assert client.post(url+'/confirm',headers=h).status_code==409
    assert client.delete(url+f'/items/{item["id"]}',headers=h).status_code==404
    again=client.post(url+'/items',headers=h,json={'product_id':products[0]['id'],'quantity':1,'price':0})
    assert again.status_code==201 and again.json()['row_number']==1 and again.json()['total']==0
    actions=[x['action'] for x in client.get(url+'/audit',headers=h).json()]
    assert actions.count('add_item')==2 and actions.count('delete_item')==2


def test_row_changes_validate_ownership_and_status(client,monkeypatch):
    from core.repositories import DocumentRepo
    _,ha=signup(client);pa=catalog(client,ha);doc=upload(client,ha)
    result=process(client,ha,doc,monkeypatch)
    item=result['items'][0]
    url=f'/api/documents/{doc["id"]}'
    _,hb=signup(client,'shop_b');pb=catalog(client,hb);other=upload(client,hb)
    process(client,hb,other,monkeypatch)
    payload={'product_id':pa[0]['id'],'quantity':1,'price':1}
    assert client.post(url+'/items',headers=hb,json=payload).status_code==404
    assert client.delete(url+f'/items/{item["id"]}',headers=hb).status_code==404
    assert client.delete(f'/api/documents/{other["id"]}/items/{item["id"]}',headers=hb).status_code==404
    assert client.post(url+'/items',headers=ha,json={**payload,'product_id':pb[0]['id']}).status_code==404
    for invalid in ({'quantity':0},{'quantity':-1},{'quantity':'NaN'},{'price':-1},{'price':'Infinity'},{'quantity':'0.0001'},{'price':'1.001'},{'product_id':None}):
        assert client.post(url+'/items',headers=ha,json={**payload,**invalid}).status_code==422
    for status in ('received','processing','retry_pending','sent_to_1c','completed','error'):
        client.portal.call(DocumentRepo.update_status,doc['id'],status)
        assert client.post(url+'/items',headers=ha,json=payload).status_code==409
        assert client.delete(url+f'/items/{item["id"]}',headers=ha).status_code==409
    assert len(client.get(url,headers=ha).json()['items'])==1
