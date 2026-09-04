"""
excel/exporter.py

Формирует Excel-файл из списка товаров в формате, готовом для
импорта в 1С:Предприятие. Использует openpyxl для полного контроля
над форматированием (ширина колонок, стили заголовка, числовые форматы).
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config.config import TABLE_COLUMNS
from models.product import Product

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_THIN_BORDER = Border(*(Side(style="thin", color="B7B7B7") for _ in range(4)))
_COLUMN_WIDTHS = (45, 16, 18, 12, 12, 14, 10)


class ExcelExporter:
    """Отвечает за создание итогового .xlsx файла с товарами накладной."""

    def __init__(self, columns: tuple[str, ...] = TABLE_COLUMNS) -> None:
        self._columns = columns

    def export(self, products: list[Product], output_path: Path) -> Path:
        """Создает Excel-файл со списком товаров и сохраняет его на диск."""
        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Накладная"

        self._write_header(sheet)
        self._write_rows(sheet, products)
        self._apply_column_widths(sheet)
        sheet.freeze_panes = "A2"

        try:
            workbook.save(output_path)
            logger.info("Excel-файл сохранен: %s (%d товаров)", output_path, len(products))
        except OSError as exc:
            logger.error("Не удалось сохранить Excel-файл %s: %s", output_path, exc)
            raise

        return output_path

    def _write_header(self, sheet: Worksheet) -> None:
        for col_index, title in enumerate(self._columns, start=1):
            cell = sheet.cell(row=1, column=col_index, value=title)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _THIN_BORDER
        sheet.row_dimensions[1].height = 22

    def _write_rows(self, sheet: Worksheet, products: list[Product]) -> None:
        for row_offset, product in enumerate(products, start=2):
            values = product.to_row()
            for col_index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_offset, column=col_index, value=self._coerce_value(col_index, value))
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(
                    horizontal="left" if col_index == 1 else "center",
                    vertical="center",
                    wrap_text=(col_index == 1),
                )

    @staticmethod
    def _coerce_value(col_index: int, value: str):
        """Пытается привести числовые колонки (кол-во, цена, сумма) к float."""
        numeric_columns = {4, 5, 6}  # Количество, Цена, Сумма
        if col_index in numeric_columns and value:
            normalized = value.replace(" ", "").replace(",", ".")
            try:
                return float(normalized)
            except ValueError:
                return value
        return value

    def _apply_column_widths(self, sheet: Worksheet) -> None:
        for index, width in enumerate(_COLUMN_WIDTHS[: len(self._columns)], start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
