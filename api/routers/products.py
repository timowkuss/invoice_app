import asyncio
import csv
import io
import zipfile
from decimal import Decimal, InvalidOperation
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook, Workbook
from core.models.product import Product
from core.repositories import ProductRepo, ProductAliasRepo, AuditLogRepo
from core.models.audit_log import AuditLog
from core.database import get_transaction
from core.uploads import read_upload
from api.routers.auth import require_store

router=APIRouter()
HEADERS={'код':'one_c_id','код 1с':'one_c_id','one_c_id':'one_c_id','наименование':'name','название':'name','name':'name',
         'артикул':'article','article':'article','штрихкод':'barcode','barcode':'barcode','единица':'unit','ед.':'unit','unit':'unit','цена':'price','price':'price'}

def parse_catalog(data,filename,store_id):
    if filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist())>50*1024*1024:
                raise ValueError('Распакованный Excel превышает 50 МБ')
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        rows=book.active.iter_rows(values_only=True)
    elif filename.lower().endswith('.csv'):
        try:
            text=data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text=data.decode('cp1251')
        delimiter=';' if text.splitlines()[0].count(';')>text.splitlines()[0].count(',') else ','
        rows=csv.reader(io.StringIO(text),delimiter=delimiter)
        book=None
    else:
        raise ValueError('Загрузите XLSX или CSV')
    products=[]
    seen=set()
    try:
        headers=[HEADERS.get(str(v or '').strip().lower()) for v in next(rows,[])]
        if 'one_c_id' not in headers or 'name' not in headers:
            raise ValueError('Нужны колонки «Код 1С» и «Название». Скачайте шаблон.')
        if len([h for h in headers if h])!=len(set(h for h in headers if h)):
            raise ValueError('Повторяются заголовки колонок')
        for line,row in enumerate(rows,2):
            if line>10001:
                raise ValueError('За один импорт можно загрузить не более 10 000 строк')
            if not any(v is not None and str(v).strip() for v in row):
                continue
            values={key: str(value).strip() if value is not None else '' for key,value in zip(headers,row) if key}
            if not values.get('one_c_id') or not values.get('name'):
                raise ValueError(f'Строка {line}: пустой код или название')
            if values['one_c_id'] in seen:
                raise ValueError(f'Строка {line}: код 1С повторяется')
            for key,value in values.items():
                if value.startswith('='):
                    raise ValueError(f'Строка {line}: замените формулы значениями')
                limit=500 if key=='name' else (50 if key=='unit' else 100)
                if len(value)>limit:
                    raise ValueError(f'Строка {line}: слишком длинное поле {key}')
            seen.add(values['one_c_id'])
            price=values.get('price')
            values['price']=Decimal(price.replace(',','.').replace(' ','')) if price else None
            if values['price'] is not None and (not values['price'].is_finite() or not 0<=values['price']<=Decimal('9999999999.99')):
                raise ValueError(f'Строка {line}: неверная цена')
            products.append(Product(store_id=store_id,**values))
        if not products:
            raise ValueError('В файле нет товаров')
        return products
    finally:
        if book:
            book.close()

@router.get('/template')
async def template(user=Depends(require_store)):
    book=Workbook()
    book.active.append(['Код 1С','Название','Артикул','Штрихкод','Единица','Цена'])
    book.active.append(['000001','Шадринское Молоко 1л 5%','','','шт',100])
    for col in 'ABCD':
        book.active[f'{col}2'].number_format='@'
    buffer=io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="catalog-template.xlsx"'})

@router.post('/import')
async def import_catalog(file:UploadFile=File(...),user=Depends(require_store)):
    if not user.is_store_admin:
        raise HTTPException(403,'Каталог может загружать администратор магазина')
    data=await read_upload(file,10*1024*1024)
    try:
        products=await asyncio.to_thread(parse_catalog,data,file.filename or '',user.store_id)
    except (ValueError,InvalidOperation,zipfile.BadZipFile,KeyError):
        raise HTTPException(422,'Некорректный каталог. Проверьте коды, названия и цены; используйте шаблон.') from None
    async with get_transaction() as conn:
        await conn.execute('UPDATE stores SET updated_at=updated_at WHERE id=$1',user.store_id)
        count=await ProductRepo.bulk_upsert(products)
        await AuditLogRepo.create(AuditLog(store_id=user.store_id,user_id=user.id,action='import_catalog',entity_type='product',new_value={'count':count}))
    return {'imported':count}

@router.get('/')
async def list_products(search:str|None=Query(None,max_length=200),limit:int=Query(50,ge=1,le=200),offset:int=Query(0,ge=0),user=Depends(require_store)):
    products=await ProductRepo.search(user.store_id,search,limit) if search else await ProductRepo.list_by_store(user.store_id,limit,offset)
    return {'items':[p.to_dict() for p in products],'total':await ProductRepo.count_by_store(user.store_id)}

@router.get('/{product_id}')
async def get_product(product_id:int,user=Depends(require_store)):
    p=await ProductRepo.get_by_id(product_id)
    if not p or p.store_id!=user.store_id:
        raise HTTPException(404,'Товар не найден')
    return {'product':p.to_dict(),'aliases':[a.to_dict() for a in await ProductAliasRepo.list_by_product(product_id)]}

@router.get('/{product_id}/aliases')
async def aliases(product_id:int,user=Depends(require_store)):
    return (await get_product(product_id,user))['aliases']
