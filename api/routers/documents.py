from __future__ import annotations
import asyncio
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field, ConfigDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from core.config import UPLOAD_DIR
from core.database import get_connection, get_transaction
from core.repositories import DocumentRepo, DocumentItemRepo, ProductRepo, ProductAliasRepo, AuditLogRepo
from core.models.document_item import DocumentItem
from core.models.product_alias import ProductAlias
from core.matching.service import normalize_text, MatchingService
from core.uploads import read_upload, document_path
from core.document_service import owned_document,lock_document,audit,save_document,queue_document
from api.routers.auth import require_store

router=APIRouter()

@router.post('/upload',status_code=201)
async def upload_document(file:UploadFile=File(...),supplier:str=Form('',max_length=500),user=Depends(require_store)):
    data=await read_upload(file)
    doc=await save_document(data,file.filename,user,supplier,upload_dir=UPLOAD_DIR)
    return public_document(doc)

def public_document(doc):
    result=doc.to_dict()
    result.pop('original_file_url',None)
    result.pop('original_file_hash',None)
    result['file_type']=Path(doc.original_file_url).suffix.lower()
    return result

@router.post('/{doc_id}/process',status_code=202)
@router.post('/{doc_id}/retry',status_code=202)
async def process_document(doc_id:int,user=Depends(require_store)):
    return await queue_document(doc_id,user)

@router.get('/')
async def list_documents(status:str|None=None,limit:int=Query(50,ge=1,le=200),offset:int=Query(0,ge=0),user=Depends(require_store)):
    docs=await DocumentRepo.list_by_store(user.store_id,status=status,limit=limit,offset=offset)
    return {'items':[public_document(doc) for doc in docs],'total':await DocumentRepo.count_by_store(user.store_id,status)}

@router.get('/{doc_id}')
async def get_document(doc_id:int,user=Depends(require_store)):
    doc=await owned_document(doc_id,user)
    items=await DocumentItemRepo.list_by_document(doc_id)
    return {'document':public_document(doc),'items':[item.to_dict() for item in items]}

@router.get('/{doc_id}/image')
async def image(doc_id:int,user=Depends(require_store)):
    doc=await owned_document(doc_id,user)
    path=document_path(doc.original_file_url)
    return FileResponse(path,headers={'Cache-Control':'private, no-store'})

@router.get('/{doc_id}/items/{item_id}/suggestions')
async def suggestions(doc_id:int,item_id:int,user=Depends(require_store)):
    await owned_document(doc_id,user)
    item=await DocumentItemRepo.get_by_id(item_id)
    if not item or item.document_id!=doc_id:
        raise HTTPException(404,'Строка не найдена')
    match=await MatchingService().match_product(user.store_id,item.ocr_text,item.article,item.barcode)
    return {'items':[p.to_dict() for p in match.alternatives],'confidence':match.confidence}

class ItemUpdate(BaseModel):
    model_config=ConfigDict(extra='forbid')
    product_id:int|None=Field(default=None,gt=0)
    quantity:Decimal|None=Field(default=None,gt=0,le=9999999,max_digits=10,decimal_places=3,allow_inf_nan=False)
    price:Decimal|None=Field(default=None,ge=0,le=999999999,max_digits=12,decimal_places=2,allow_inf_nan=False)
    create_alias:bool=False

class ItemCreate(BaseModel):
    model_config=ConfigDict(extra='forbid')
    product_id:int=Field(gt=0)
    quantity:Decimal=Field(gt=0,le=9999999,max_digits=10,decimal_places=3,allow_inf_nan=False)
    price:Decimal=Field(ge=0,le=999999999,max_digits=12,decimal_places=2,allow_inf_nan=False)

def require_editable(doc):
    if doc.status not in ('needs_review','recognized','confirmed'):
        raise HTTPException(409,'Редактирование доступно после распознавания')

async def refresh_items(doc_id,conn):
    items=await DocumentItemRepo.list_by_document(doc_id)
    good=sum(i.match_status in ('matched','manual') for i in items)
    await DocumentRepo.update_counts(doc_id,len(items),good,len(items)-good)
    await DocumentRepo.update_status(doc_id,'needs_review')
    await conn.execute('UPDATE documents SET confirmed_at=NULL WHERE id=$1',doc_id)

@router.post('/{doc_id}/items',status_code=201)
async def add_item(doc_id:int,req:ItemCreate,user=Depends(require_store)):
    async with get_transaction() as conn:
        doc=await lock_document(doc_id,user,conn)
        require_editable(doc)
        product=await ProductRepo.get_by_id(req.product_id)
        if not product or product.store_id!=user.store_id or not product.is_active:
            raise HTTPException(404,'Товар не найден в каталоге магазина')
        items=await DocumentItemRepo.list_by_document(doc_id)
        if len(items)>=2000:
            raise HTTPException(409,'В накладной не может быть больше 2000 строк')
        item=await DocumentItemRepo.create(DocumentItem(
            document_id=doc_id,row_number=max((i.row_number for i in items),default=0)+1,
            product_id=product.id,product_name=product.name,article=product.article or '',
            barcode=product.barcode or '',unit=product.unit or '',quantity=req.quantity,price=req.price,
            total=(req.quantity*req.price).quantize(Decimal('.01'),rounding=ROUND_HALF_UP),
            match_status='manual',confidence=Decimal('100')))
        await refresh_items(doc_id,conn)
        await audit(user,doc_id,'add_item',new=item.to_dict())
    return item.to_dict()

@router.delete('/{doc_id}/items/{item_id}')
async def delete_item(doc_id:int,item_id:int,user=Depends(require_store)):
    async with get_transaction() as conn:
        doc=await lock_document(doc_id,user,conn)
        require_editable(doc)
        item=await DocumentItemRepo.get_by_id(item_id)
        if not item or item.document_id!=doc_id:
            raise HTTPException(404,'Строка не найдена')
        await conn.execute('DELETE FROM document_items WHERE id=$1 AND document_id=$2',item_id,doc_id)
        items=await DocumentItemRepo.list_by_document(doc_id)
        for number,row in enumerate(items,1):
            await conn.execute('UPDATE document_items SET row_number=$1 WHERE id=$2',number,row.id)
        await refresh_items(doc_id,conn)
        await audit(user,doc_id,'delete_item',old=item.to_dict())
    return {'status':'deleted'}

@router.put('/{doc_id}/items/{item_id}')
async def update_item(doc_id:int,item_id:int,req:ItemUpdate,user=Depends(require_store)):
    async with get_transaction() as conn:
        doc=await lock_document(doc_id,user,conn)
        if doc.status not in ('needs_review','recognized','confirmed'):
            raise HTTPException(409,'Дождитесь завершения распознавания')
        item=await DocumentItemRepo.get_by_id(item_id)
        if not item or item.document_id!=doc_id:
            raise HTTPException(404,'Строка не найдена')
        old=item.to_dict()
        if req.product_id:
            product=await ProductRepo.get_by_id(req.product_id)
            if not product or product.store_id!=user.store_id or not product.is_active:
                raise HTTPException(404,'Товар не найден в каталоге магазина')
            item.product_id=product.id
            item.product_name=product.name
            item.article=product.article or ''
            item.barcode=product.barcode or ''
            item.unit=product.unit or ''
            item.match_status='manual'
            item.confidence=Decimal('100')
        if req.quantity is not None:
            item.quantity=req.quantity
        if req.price is not None:
            item.price=req.price
        if item.quantity is not None and item.price is not None:
            item.total=(Decimal(str(item.quantity))*Decimal(str(item.price))).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        item=await DocumentItemRepo.update(item)
        if req.create_alias and req.product_id and item.ocr_text:
            # The user explicitly teaches this store; other stores never see the alias.
            normalized=normalize_text(item.ocr_text)
            await conn.execute('UPDATE stores SET updated_at=updated_at WHERE id=$1',user.store_id)
            await conn.execute('DELETE FROM product_aliases WHERE store_id=$1 AND normalized_text=$2',user.store_id,normalized)
            await ProductAliasRepo.create(ProductAlias(store_id=user.store_id,product_id=req.product_id,ocr_text=item.ocr_text,normalized_text=normalized,confidence=Decimal('100'),created_by=user.id))
        await refresh_items(doc_id,conn)
        await audit(user,doc_id,'update_item',old,item.to_dict())
    return item.to_dict()

async def reviewed_items(doc_id,store_id):
    items=await DocumentItemRepo.list_by_document(doc_id)
    if not items:
        raise HTTPException(409,'В документе нет товаров')
    result=[]
    for item in items:
        p=await ProductRepo.get_by_id(item.product_id) if item.product_id else None
        if not p or p.store_id!=store_id or not p.is_active or item.match_status not in ('matched','manual'):
            raise HTTPException(409,f'Строка {item.row_number}: выберите товар из каталога')
        if item.quantity is None or item.price is None or item.total is None:
            raise HTTPException(409,f'Строка {item.row_number}: заполните количество и цену')
        quantity,price,total=map(lambda x:Decimal(str(x)),(item.quantity,item.price,item.total))
        if not all(v.is_finite() for v in (quantity,price,total)) or quantity<=0 or price<0 or abs(quantity*price-total)>Decimal('.02'):
            raise HTTPException(409,f'Строка {item.row_number}: проверьте количество, цену и сумму')
        if item.product_name!=p.name or (item.unit or '')!=(p.unit or ''):
            raise HTTPException(409,f'Строка {item.row_number}: каталог изменился, выберите товар заново')
        result.append((item,p))
    return result

@router.post('/{doc_id}/confirm')
async def confirm(doc_id:int,user=Depends(require_store)):
    async with get_transaction() as conn:
        doc=await lock_document(doc_id,user,conn)
        if doc.status not in ('recognized','needs_review','confirmed'):
            raise HTTPException(409,'Документ ещё не готов к проверке')
        await reviewed_items(doc_id,user.store_id)
        await DocumentRepo.update_status(doc_id,'confirmed')
        await audit(user,doc_id,'confirm')
    return {'status':'confirmed'}

def workbook_bytes(rows):
    book=Workbook()
    sheet=book.active
    sheet.title='Накладная'
    sheet.append(['Код 1С','Название','Артикул','Штрихкод','Количество','Цена','Сумма','Единица'])
    for item,p in rows:
        sheet.append([p.one_c_id,p.name,p.article or '',p.barcode or '',float(item.quantity),float(item.price),float(item.total),p.unit or ''])
        for col in (1,2,3,4,8):
            cell=sheet.cell(sheet.max_row,col)
            cell.data_type='s'  # Untrusted OCR/catalog values must never become formulas.
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='194C40')
    for col,width in zip('ABCDEFGH',(24,55,20,22,15,15,15,14)):
        sheet.column_dimensions[col].width=width
    for row in sheet.iter_rows(min_row=2):
        row[4].number_format='0.000'
        row[5].number_format=row[6].number_format='0.00'
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    buffer=BytesIO()
    book.save(buffer)
    return buffer.getvalue()

@router.get('/{doc_id}/export')
async def export(doc_id:int,user=Depends(require_store)):
    async with get_transaction() as conn:
        doc=await lock_document(doc_id,user,conn)
        if doc.status!='confirmed':
            raise HTTPException(409,'Сначала проверьте и подтвердите накладную')
        rows=await reviewed_items(doc_id,user.store_id)
        content=await asyncio.to_thread(workbook_bytes,rows)
        await audit(user,doc_id,'export_excel')
    return StreamingResponse(BytesIO(content),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="invoice-{doc_id}.xlsx"','Cache-Control':'private, no-store'})

@router.get('/{doc_id}/audit')
async def history(doc_id:int,user=Depends(require_store)):
    await owned_document(doc_id,user)
    return [entry.to_dict() for entry in await AuditLogRepo.list_by_document(doc_id)]
